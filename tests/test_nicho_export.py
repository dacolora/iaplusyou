import io


SUB_A = {"id": 11, "nombre": "Melissa / La que regala", "deseo": "Quiero regalar bien", "base": "emocion", "estado": "aprobado",
         "demografia": "Mujer 30-50", "edad_rango": "30-50", "emocion": "Presión",
         "identidad": {"quiere_que_vean": "detallista", "cree_de_si": "generosa", "quiere_lograr": "acertar"},
         "soluciones_previas": [{"que": "Tarjetas", "por_que_fallo": ["impersonal", "sin valor"]}, {"que": "Medias", "por_que_fallo": ["genéricas"]}],
         "situaciones": ["Compra a última hora", "Centro comercial"], "comportamiento": "Va a lo seguro",
         "conciencia": {"nivel": "consciente_del_problema", "detalle": "Sabe que falla"}, "encaje_producto": "Pantuflas con soporte",
         "tono": "Cálido", "palabras_clave": ["regalo"], "evidencia": [{"comentario_id": 1, "cita": "nunca sé qué regalar"}], "sin_evidencia": False}
SUB_B = dict(SUB_A, id=12, nombre="Descartado", estado="descartado")
NUCLEOS = [{"id": 1, "nombre": "Regalo con cabeza", "deseo": "Quiero regalar bien", "resumen": "Quienes regalan", "subs": [SUB_A, SUB_B]},
           {"id": 2, "nombre": "Vacío", "deseo": "Quiero x", "resumen": "", "subs": []}]
ESTUDIO = {"id": 5, "nombre": "Pantuflas / regalo", "producto": "HappyFlops", "tema": "regalos", "idioma": "es", "generacion": 2,
           "extra": {"ultima_generacion": {"fecha": "2026-09-18T10:00:00"}}}


def test_valor_y_subs_exportables():
    from nicho import exportar
    assert [(n["id"], s["id"]) for n, s in exportar.subs_exportables(NUCLEOS)] == [(1, 11)]
    assert exportar.valor(SUB_A, "identidad.cree_de_si") == "generosa"
    assert exportar.valor(SUB_A, "soluciones_previas.que") == "1. Tarjetas\n2. Medias"
    assert exportar.valor(SUB_A, "soluciones_previas.por_que_fallo") == "1. impersonal; sin valor\n2. genéricas"
    assert exportar.valor(SUB_A, "situaciones") == "Compra a última hora\nCentro comercial"
    assert exportar.valor(SUB_A, "conciencia") == "consciente del problema — Sabe que falla"
    assert exportar.valor(SUB_A, "evidencia", urls={1: "https://r.com/1"}) == "«nunca sé qué regalar» (https://r.com/1)"
    assert exportar.valor(SUB_A, "evidencia") == "«nunca sé qué regalar»"
    assert exportar.valor({}, "tono") == ""
    assert [r for _, r in exportar.FILAS_HOJA][:3] == ["Personas", "Deseo (frase de cabecera)", "Demographics (ASL)"]


def test_markdown():
    from nicho import exportar
    md = exportar.markdown(ESTUDIO, NUCLEOS, urls={1: "https://r.com/1"})
    for frag in ("# Avatares: Pantuflas / regalo", "HappyFlops", "## Núcleo 1: Regalo con cabeza", "**Deseo:** Quiero regalar bien",
                 "### Sub-avatar 1.1: Melissa / La que regala", "emoción · aprobado", "**Beliefs about self:** generosa",
                 "1. Tarjetas", "«nunca sé qué regalar» (https://r.com/1)", "## Núcleo 2: Vacío"):
        assert frag in md, frag
    assert "Descartado" not in md
    assert exportar.nombre_archivo(ESTUDIO, "md") == "avatares_pantuflas-regalo_5.md"


def test_excel():
    from openpyxl import load_workbook
    from nicho import exportar
    wb = load_workbook(io.BytesIO(exportar.excel(ESTUDIO, NUCLEOS)))
    ws = wb["Personas"]
    assert ws["A1"].value == "Personas" and ws["B1"].value == "Melissa / La que regala" and ws["C1"].value is None
    assert ws["A3"].value == "Demographics (ASL)" and ws["B3"].value == "Mujer 30-50"
    assert ws["A5"].value == "Beliefs about self" and ws["B5"].value == "generosa"
    assert ws["A8"].value.startswith("What are other solutions") and ws["B8"].value == "1. Tarjetas\n2. Medias"
    # Verificar que la exportación neutraliza celdas que empiezan como fórmula
    SUB_F = dict(SUB_A, id=13, nombre="=SUM(A1)", tono="+peligro", estado="propuesto")
    wb2 = load_workbook(io.BytesIO(exportar.excel(ESTUDIO, [{"id": 9, "nombre": "N", "deseo": "d", "resumen": "", "subs": [SUB_F]}])))
    ws2 = wb2["Personas"]
    assert ws2["B1"].value == "'=SUM(A1)" and ws2["B14"].value == "'+peligro"
    assert ws["B16"].value == "«nunca sé qué regalar»" and ws.freeze_panes == "B2"


def test_urls_comentarios(base_temporal):
    from nicho import datos
    eid = datos.crear_estudio("acme", "X")
    datos.agregar_comentarios("acme", eid, "texto", [{"fuente_id": "a", "texto": "con link", "url": "https://r.com/1"},
                                                     {"fuente_id": "b", "texto": "sin link"}])
    urls = datos.urls_comentarios("acme", eid)
    assert set(urls.values()) == {None, "https://r.com/1"}
    assert datos.urls_comentarios("otro", eid) == {}
