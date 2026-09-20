"""
Exportación de un estudio (spec §6): `.md` como el doc "Desire-Based Core
Avatar" y `.xlsx` con la plantilla de la hoja "Personas" del cliente
(columna A los rótulos tal cual la hoja, una columna por sub-avatar; los
descartados no salen). Solo openpyxl; nada de Flask ni de base de datos.
"""
import io
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# (clave, rótulo). Las primeras filas son la hoja "Personas" con sus rótulos
# exactos; después van los campos del doc del detergente que la hoja no tiene.
FILAS_HOJA = [
    ("nombre", "Personas"),
    ("deseo", "Deseo (frase de cabecera)"),
    ("demografia", "Demographics (ASL)"),
    ("identidad.quiere_que_vean", "What are some of the characteristics your prospect wants others to see in them?"),
    ("identidad.cree_de_si", "Beliefs about self"),
    ("identidad.quiere_lograr", "What does the prospect want to achieve in society?"),
    ("encaje_producto", "How does your product help them achieve that status/characteristics?"),
    ("soluciones_previas.que", "What are other solutions they have tried and failed at?"),
    ("soluciones_previas.por_que_fallo", "Reason for failure with those solutions"),
    ("situaciones", "Su día a día (situaciones)"),
    ("conciencia", "Nivel de conciencia"),
    ("emocion", "Emoción"),
    ("comportamiento", "Comportamiento"),
    ("tono", "Tono de voz"),
    ("palabras_clave", "Palabras clave"),
    ("evidencia", "Citas textuales"),
]
BASES_NOMBRE = {"emocion": "emoción", "experiencia_producto": "experiencia con el producto"}

# Igual que tablero._celda: un texto que empieza por =, +, -, @, tab o CR lo
# ejecutaría Excel/Sheets como fórmula al abrir el archivo. Los campos vienen de
# comentarios ajenos y de Claude, así que se neutralizan con un apóstrofo.
_INICIOS_FORMULA = ("=", "+", "-", "@", "\t", "\r")


def _celda(texto):
    t = "" if texto is None else str(texto)
    return "'" + t if t.startswith(_INICIOS_FORMULA) else t


def subs_exportables(nucleos):
    return [(n, s) for n in (nucleos or []) for s in (n.get("subs") or []) if s.get("estado") != "descartado"]


def _cita(e, urls):
    url = (urls or {}).get(e.get("comentario_id"))
    return f"«{e.get('cita', '')}»" + (f" ({url})" if url else "")


def valor(sub, clave, urls=None):
    """Texto plano de una fila de la hoja para un sub-avatar."""
    sub = sub or {}
    if clave == "conciencia":
        c = sub.get("conciencia") or {}
        nivel = (c.get("nivel") or "").replace("_", " ")
        return " — ".join(x for x in (nivel, c.get("detalle") or "") if x)
    if clave == "evidencia":
        return "\n".join(_cita(e, urls) for e in sub.get("evidencia") or [])
    if clave == "soluciones_previas.que":
        return "\n".join(f"{i + 1}. {s.get('que') or ''}" for i, s in enumerate(sub.get("soluciones_previas") or []))
    if clave == "soluciones_previas.por_que_fallo":
        return "\n".join(f"{i + 1}. " + "; ".join(s.get("por_que_fallo") or [])
                         for i, s in enumerate(sub.get("soluciones_previas") or []))
    if clave.startswith("identidad."):
        return (sub.get("identidad") or {}).get(clave.split(".", 1)[1]) or ""
    v = sub.get(clave)
    if isinstance(v, list):
        return "\n".join(str(x) for x in v)
    return "" if v is None else str(v)


def _slug(texto):
    texto = unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9]+", "-", texto.strip()).strip("-").lower()[:40] or "estudio"


def nombre_archivo(estudio, ext):
    return f"avatares_{_slug(estudio.get('nombre'))}_{estudio.get('id')}.{ext}"


def markdown(estudio, nucleos, urls=None):
    fecha = ((estudio.get("extra") or {}).get("ultima_generacion") or {}).get("fecha") or ""
    lineas = [f"# Avatares: {estudio.get('nombre') or ''}", "",
              f"Producto: {estudio.get('producto') or '—'} · Tema: {estudio.get('tema') or '—'} · Idioma: {estudio.get('idioma') or 'es'}"
              f" · Generación {estudio.get('generacion') or 0}" + (f" · {fecha}" if fecha else ""), ""]
    for i, n in enumerate(nucleos or [], start=1):
        lineas += [f"## Núcleo {i}: {n.get('nombre') or ''}", "", f"**Deseo:** {n.get('deseo') or ''}", ""]
        if n.get("resumen"):
            lineas += [n["resumen"], ""]
        subs = [s for s in (n.get("subs") or []) if s.get("estado") != "descartado"]
        for j, s in enumerate(subs, start=1):
            lineas += [f"### Sub-avatar {i}.{j}: {s.get('nombre') or ''} "
                       f"({BASES_NOMBRE.get(s.get('base'), s.get('base') or '')} · {s.get('estado') or 'propuesto'})", ""]
            for clave, rotulo in FILAS_HOJA[1:]:
                v = valor(s, clave, urls)
                if not v:
                    continue
                if "\n" in v:
                    lineas.append(f"- **{rotulo}:**")
                    lineas += [f"  - {x}" for x in v.split("\n")]
                else:
                    lineas.append(f"- **{rotulo}:** {v}")
            lineas.append("")
    return "\n".join(lineas).rstrip() + "\n"


def excel(estudio, nucleos, urls=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Personas"
    negrita = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="3B6FD9")
    ajuste = Alignment(wrap_text=True, vertical="top")
    for fila, (_, rotulo) in enumerate(FILAS_HOJA, start=1):
        celda = ws.cell(row=fila, column=1, value=rotulo)
        celda.font, celda.fill, celda.alignment = negrita, relleno, ajuste
    for col, (_, s) in enumerate(subs_exportables(nucleos), start=2):
        for fila, (clave, _) in enumerate(FILAS_HOJA, start=1):
            celda = ws.cell(row=fila, column=col, value=_celda(valor(s, clave, urls)) or None)
            celda.alignment = ajuste
            if fila == 1:
                celda.font = Font(bold=True)
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 45
    ws.column_dimensions["A"].width = 38
    ws.freeze_panes = "B2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
