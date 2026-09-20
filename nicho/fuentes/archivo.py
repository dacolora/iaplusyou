"""
Fuente `csv` (spec §3.2): un archivo .csv (delimitador detectado con
csv.Sniffer) o .xlsx (primera hoja, openpyxl) con fila de encabezados. La
columna de texto se reconoce por nombre (COLUMNAS_TEXTO) o, si no hay, es la
de mayor largo medio; url, puntuación y fecha son opcionales por nombre.
Máximo MAX_BYTES y MAX_FILAS. Corre en la ruta, sin worker.
"""
import csv
import io
import os

from nicho.fuentes.base import ErrorFuente, Fuente, normalizar_comentario

MAX_BYTES = 5 * 1024 * 1024
MAX_FILAS = 5000
COLUMNAS_TEXTO = ("texto", "comentario", "comment", "review", "body", "content", "text", "reseña", "resena",
                  "mensaje", "opinion", "opinión")
COLUMNAS_URL = ("url", "link", "enlace")
COLUMNAS_PUNTUACION = ("score", "likes", "votos", "puntuacion", "puntuación", "rating", "upvotes")
COLUMNAS_FECHA = ("fecha", "date", "created", "published")
_MIN_LARGO_MEDIO = 10        # menos que esto no parece una columna de comentarios


def _decodificar(contenido):
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return contenido.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ErrorFuente("No pude leer el archivo como texto (¿es un CSV?).")


def _sin_vacias(filas):
    return [f for f in filas if any(str(c or "").strip() for c in f)]


def _filas_csv(contenido):
    texto = _decodificar(contenido)
    try:
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t|")
    except csv.Error:
        dialecto = csv.excel
    filas = _sin_vacias(list(csv.reader(io.StringIO(texto), dialecto)))
    if not filas:
        raise ErrorFuente("El archivo está vacío.")
    return [str(c or "").strip() for c in filas[0]], filas[1:]


def _filas_xlsx(contenido):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — openpyxl lanza de todo con un archivo dañado
        raise ErrorFuente("No pude abrir el Excel (¿está dañado o protegido?).") from None
    try:
        hoja = wb.worksheets[0]
        filas = []
        for cruda in hoja.iter_rows(values_only=True):
            fila = ["" if v is None else str(v) for v in cruda]
            if not any(c.strip() for c in fila):
                continue
            filas.append(fila)
            # Corta apenas se pasa (encabezado + MAX_FILAS): una hoja enorme
            # nunca se termina de traer a memoria solo para descartarla después.
            if len(filas) > MAX_FILAS + 1:
                raise ErrorFuente(f"El archivo tiene más de {MAX_FILAS} filas; pártelo.")
    finally:
        wb.close()
    if not filas:
        raise ErrorFuente("El Excel está vacío.")
    return [c.strip() for c in filas[0]], filas[1:]


def _indice_por_nombre(encabezados, nombres):
    bajos = [e.strip().lower() for e in encabezados]
    for n in nombres:
        if n in bajos:
            return bajos.index(n)
    return None


def detectar_columnas(encabezados, filas):
    texto = _indice_por_nombre(encabezados, COLUMNAS_TEXTO)
    if texto is None:
        mejor, mejor_largo = None, 0.0
        for i in range(len(encabezados)):
            largos = [len(str(f[i])) for f in filas[:200] if i < len(f) and str(f[i] or "").strip()]
            media = sum(largos) / len(largos) if largos else 0.0
            if media > mejor_largo:
                mejor, mejor_largo = i, media
        if mejor is None or mejor_largo < _MIN_LARGO_MEDIO:
            raise ErrorFuente("No encontré una columna con comentarios (ponle «texto» o «comentario» de encabezado).")
        texto = mejor
    return {"texto": texto, "url": _indice_por_nombre(encabezados, COLUMNAS_URL),
            "puntuacion": _indice_por_nombre(encabezados, COLUMNAS_PUNTUACION),
            "fecha": _indice_por_nombre(encabezados, COLUMNAS_FECHA)}


def leer_archivo(nombre, contenido):
    """-> comentarios normalizados. ErrorFuente si el archivo no sirve."""
    contenido = contenido or b""
    if len(contenido) > MAX_BYTES:
        raise ErrorFuente("El archivo pesa más de 5 MB; pártelo.")
    ext = os.path.splitext(nombre or "")[1].lower()
    if ext == ".csv":
        encabezados, filas = _filas_csv(contenido)
    elif ext in (".xlsx", ".xlsm"):
        encabezados, filas = _filas_xlsx(contenido)
    else:
        raise ErrorFuente("Solo acepto archivos .csv o .xlsx.")
    if len(filas) > MAX_FILAS:
        raise ErrorFuente(f"El archivo tiene más de {MAX_FILAS} filas; pártelo.")
    cols = detectar_columnas(encabezados, filas)

    def celda(f, i):
        return f[i] if i is not None and i < len(f) else None

    salida = []
    for f in filas:
        c = normalizar_comentario({"texto": celda(f, cols["texto"]), "url": celda(f, cols["url"]),
                                   "puntuacion": celda(f, cols["puntuacion"]), "fecha": celda(f, cols["fecha"])})
        if c:
            salida.append(c)
    return salida


class FuenteArchivo(Fuente):
    tipo = "csv"

    def recolectar(self, params, avanzar=None):
        params = dict(params or {})
        yield from leer_archivo(params.get("nombre"), params.get("contenido"))
